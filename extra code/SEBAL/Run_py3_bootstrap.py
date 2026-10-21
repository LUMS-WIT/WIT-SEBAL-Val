"""Run regional calibration ensembles independently of WIT-SEBAL-Val/main.py.

Run with --help. Importing this file or using --prepare-only / --preflight does
not import GDAL or execute SEBAL. Workbook column C is never used for outputs.
"""
from pathlib import Path
import argparse
import contextlib
import csv
import gzip
import json
import math
import re
import sys
import tempfile
import traceback

from openpyxl import load_workbook
from sebal_bootstrap_core import (EnsembleContext, digest, read_coefficients,
                                  read_locations, select_coefficients, write_table,
                                  coefficient_subset_diagnostics)

HERE = Path(__file__).resolve().parent
INPUTS = HERE/'ensemble_inputs'
if len(HERE.parents) >= 2:
    repo_inputs = HERE.parents[1]/'uncertainty'/'bootstrap'/'outputs'/'samples'
    if (repo_inputs/'coefficient_bank.csv').exists():
        INPUTS = repo_inputs


def scene_rows(workbook, start=2, end=None):
    book = load_workbook(workbook, data_only=True, read_only=True)
    rows = []
    try:
        general, landsat = book['General_Input'], book['Landsat_Input']
        for row in range(start, min(end or general.max_row, general.max_row)+1):
            if general.cell(row, 2).value is None:
                continue
            if int(general.cell(row, 4).value) != 1:
                raise ValueError(f'{workbook}, row {row}: only Landsat Image_Type=1 is supported.')
            scene = str(landsat.cell(row, 2).value)
            parts = scene.split('_')
            if len(parts) < 4 or not re.fullmatch(r'\d{6}', parts[2]) or not re.fullmatch(r'\d{8}', parts[3]):
                raise ValueError(f'{workbook}, row {row}: unrecognized scene name {scene}')
            date = parts[3]
            metadata = dict(workbook=str(Path(workbook).resolve()), excel_row=row, scene=scene,
                            date=f'{date[:4]}-{date[4:6]}-{date[6:8]}', swath=parts[2])
            paths = [str(general.cell(row, 2).value), str(general.cell(row, 5).value)]
            # Raster paths supplied in other sheets must exist; numeric constants
            # remain numeric. Column C of General_Input is deliberately excluded.
            for sheet in book:
                if sheet.title not in ['General_Input', 'Meteo_Input', 'Soil_Input',
                                       'Landsat_Input', 'Additional_Input']:
                    continue  # VIIRS/MODIS example cells are inactive for Landsat.
                for cells in sheet.iter_rows(min_row=row, max_row=row):
                    for column, cell in enumerate(cells, 1):
                        value = cell.value
                        if sheet.title == 'General_Input' and column == 3:
                            continue
                        if isinstance(value, str) and re.search(r'\.(tif|tiff|hdf|nc)$', value, re.I):
                            paths.append(value)
            folder = Path(paths[0])
            paths.extend(str(folder/f'{scene}{suffix}') for suffix in ['_MTL.txt', '_B2.TIF'])
            rows.append(dict(metadata=metadata, required_paths=sorted(set(paths))))
    finally:
        book.close()
    if not rows:
        raise ValueError(f'No populated scenes in requested rows of {workbook}')
    return rows


def preflight(rows, locations):
    missing = sorted({p for row in rows for p in row['required_paths'] if not Path(p).exists()})
    swaths = sorted({row['metadata']['swath'] for row in rows})
    unassigned = [s for s in swaths if not any(not p['swaths'] or s in p['swaths'] for p in locations)]
    return dict(n_scenes=len(rows), swaths=swaths, n_gpis=len(locations),
                n_missing_paths=len(missing), missing_paths=missing,
                swaths_without_sensors=unassigned, ready=not missing and not unassigned,
                note='Checks workbook paths, DEM, MTL and B2; remaining bands and GDAL are checked during model execution.')


def configuration(args, workbooks, coefficients):
    return dict(workflow='sebal_calibration_bootstrap',
                coefficients_sha256=digest(args.coefficients), locations_sha256=digest(args.locations),
                workbooks=[dict(path=str(p.resolve()), sha256=digest(p)) for p in workbooks],
                code_sha256={name: digest(HERE/name) for name in
                             ['Run_py3_bootstrap.py', 'pysebal_bootstrap.py', 'sebal_bootstrap_core.py']},
                subset_seed=args.subset_seed, n_realizations=len(coefficients),
                selected_replicates=[r['replicate'] for r in coefficients],
                central_alpha=args.central_alpha, central_beta=args.central_beta,
                reference_root=str(Path(args.reference_root).resolve()) if args.reference_root else None,
                scope='Conditional uncertainty in the existing EF calibration fit; sensor coordinates only.')


def initialize_output(output, settings, resume):
    marker = output/'run_manifest.json'
    if output.exists() and any(output.iterdir()):
        if not resume or not marker.exists() or json.loads(marker.read_text()) != settings:
            raise ValueError('Output is nonempty. Use a new folder, or --resume with identical inputs/settings/code.')
    output.mkdir(parents=True, exist_ok=True)
    marker.write_text(json.dumps(settings, indent=2)+'\n')


def combine_tables(output, completed):
    """Stream scene outputs into one set per workbook; never hold long tables in RAM."""
    groups = {}
    for metadata, folder in completed:
        groups.setdefault(metadata['workbook'], []).append(folder)
    exports = output/'by_workbook'
    exports.mkdir(exist_ok=True)
    used = set()
    for workbook, folders in groups.items():
        name = Path(workbook).stem
        if name in used:
            raise ValueError('Workbook file names must be distinct for combined exports.')
        used.add(name)
        for filename in ['point_realizations.csv.gz', 'point_summary.csv', 'downstream_convergence.csv']:
            opener = gzip.open if filename.endswith('.gz') else open
            target = exports/f'{name}_{filename}'
            # An interrupted combine cannot corrupt a previously completed export.
            temporary = exports/f'.{name}_{filename}.tmp'
            with opener(temporary, 'wt', newline='', encoding='utf-8') as dest:
                writer = None
                for folder in folders:
                    with opener(folder/filename, 'rt', newline='', encoding='utf-8') as source:
                        reader = csv.DictReader(source)
                        if writer is None:
                            writer = csv.DictWriter(dest, fieldnames=reader.fieldnames)
                            writer.writeheader()
                        if reader.fieldnames != writer.fieldnames:
                            raise ValueError('Scene output columns differ.')
                        writer.writerows(reader)
            temporary.replace(target)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--excel', action='append', type=Path,
                        help='Workbook path; repeat for both swaths. Default: both supplied workbooks.')
    parser.add_argument('--coefficients', type=Path, default=INPUTS/'coefficient_bank.csv')
    repo_locations = HERE.parents[1]/'uncertainty'/'bootstrap'/'inputs'/'sensor_locations.csv'
    parser.add_argument('--locations', type=Path,
                        default=repo_locations if repo_locations.exists() else INPUTS/'sensor_locations.csv')
    parser.add_argument('--output', type=Path,
                        default=HERE.parents[1]/'uncertainty'/'bootstrap'/'inputs'/'regional'/'new_run')
    parser.add_argument('--n-realizations', type=int, default=200,
                        help='Number of paired downstream runs; 0 selects the complete coefficient bank.')
    parser.add_argument('--subset-seed', type=int, default=20261012)
    parser.add_argument('--start-row', type=int, default=2)
    parser.add_argument('--end-row', type=int, help='Inclusive; default: all populated rows.')
    parser.add_argument('--central-alpha', type=float, default=-.089)
    parser.add_argument('--central-beta', type=float, default=.343)
    parser.add_argument('--reference-root', type=Path,
                        help='Optional saved mean-raster root with 149039/150039 subfolders; checks central values before sampling.')
    parser.add_argument('--prepare-only', action='store_true', help='Save the selected coefficients and run manifest; do not execute SEBAL.')
    parser.add_argument('--preflight', action='store_true', help='Check current workbook paths; do not execute SEBAL.')
    parser.add_argument('--resume', action='store_true', help='Skip complete scenes in an identical previous run.')
    args = parser.parse_args(argv)
    if args.start_row < 2 or (args.end_row is not None and args.end_row < args.start_row):
        parser.error('Choose Excel rows beginning at 2, with end-row >= start-row.')
    if not all(math.isfinite(x) for x in [args.central_alpha, args.central_beta]):
        parser.error('Central calibration coefficients must be finite.')
    workbooks = [p.resolve() for p in (args.excel or
                 [HERE/'InputEXCEL_LBDC_149039.xlsx', HERE/'InputEXCEL_LBDC_150039.xlsx'])]
    if len({p.stem for p in workbooks}) != len(workbooks):
        parser.error('Provide workbooks with distinct file names.')
    bank = read_coefficients(args.coefficients)
    coefficients = select_coefficients(bank,
                                       args.n_realizations or None, args.subset_seed)
    locations = read_locations(args.locations)
    rows = [row for workbook in workbooks for row in scene_rows(workbook, args.start_row, args.end_row)]
    check = preflight(rows, locations)
    if args.preflight:
        print(json.dumps(check, indent=2))
        return 0 if check['ready'] else 2
    output = args.output.resolve()
    settings = configuration(args, workbooks, coefficients)
    if not args.prepare_only and not check['ready']:
        print(f"Cannot run: {check['n_missing_paths']} required paths unavailable; "
              f"swaths without sensors: {check['swaths_without_sensors']}. Use --preflight for details.", file=sys.stderr)
        return 2
    # Import the original environment's dependencies before creating run outputs.
    if not args.prepare_only:
        import pysebal_bootstrap
    initialize_output(output, settings, args.resume)
    write_table(output/'selected_coefficients.csv', coefficients)
    write_table(output/'subset_ef_diagnostics.csv', coefficient_subset_diagnostics(bank, coefficients))
    (output/'preflight.json').write_text(json.dumps(check, indent=2)+'\n')
    if args.prepare_only:
        print(f'Saved {len(coefficients)} paired settings in {output}; SEBAL was not executed.')
        return 0
    context = EnsembleContext(coefficients=coefficients, locations=locations, output=output,
                              central_alpha=args.central_alpha, central_beta=args.central_beta,
                              reference_root=args.reference_root)
    (output/'logs').mkdir(exist_ok=True)
    (output/'scratch').mkdir(exist_ok=True)
    completed, statuses = [], []
    for i, row in enumerate(rows, 1):
        metadata = row['metadata']
        target = context.scene_directory(metadata)
        print(f"Scene {i}/{len(rows)}: {metadata['swath']}, Excel row {metadata['excel_row']}, {metadata['date']}", flush=True)
        status = dict(metadata, status='', error='', output='')
        log = output/'logs'/f"{Path(metadata['workbook']).stem}_row_{metadata['excel_row']:04d}.txt"
        try:
            if args.resume and target.exists():
                marker = target/'scene_manifest.json'
                required = ['point_realizations.csv.gz', 'point_summary.csv', 'downstream_convergence.csv']
                if not marker.exists() or not json.loads(marker.read_text()).get('complete') or not all((target/name).exists() for name in required):
                    raise ValueError(f'Existing scene output is incomplete: {target}')
                status['status'] = 'skipped_complete'
            else:
                with tempfile.TemporaryDirectory(prefix=f"{metadata['swath']}_{metadata['excel_row']}_", dir=output/'scratch') as scratch:
                    with open(log, 'w', encoding='utf-8') as stream, contextlib.redirect_stdout(stream):
                        target = pysebal_bootstrap.SEBALcode(metadata['excel_row'], metadata['workbook'],
                                                          ensemble_context=context, scratch_dir=scratch)
                status['status'] = 'complete'
            status['output'] = str(target)
            completed.append((metadata, target))
            print(f"  {status['status']}: {target}", flush=True)
        except Exception as exc:
            status.update(status='failed', error=f'{type(exc).__name__}: {exc}')
            with open(log, 'a', encoding='utf-8') as stream:
                traceback.print_exc(file=stream)
            print(f"  FAILED: {exc}; see {log}", file=sys.stderr, flush=True)
        statuses.append(status)
        write_table(output/'row_status.csv', statuses)
    combine_tables(output, completed)
    failed = sum(s['status'] == 'failed' for s in statuses)
    (output/'run_summary.json').write_text(json.dumps(dict(n_requested=len(rows), n_completed=len(completed),
                                                        n_failed=failed, all_requested_complete=failed == 0), indent=2)+'\n')
    print(f'Finished: {len(completed)} complete scenes; {failed} failed. Exports: {output/"by_workbook"}')
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(main())

# ---------------------------------------------------------------------------
# Copy-and-paste terminal commands: run from the folder containing this file.
# Use your original SEBAL/GDAL environment. Each --output must be a fresh folder,
# except when resuming a run with identical inputs, options and code.
# Paths below assume this file is in the project's "extra code/SEBAL" folder.
# For portable installations, replace output/input paths for your machine.
# ---------------------------------------------------------------------------
# Show all options.
# python Run_py3_bootstrap.py --help
#
# Preflight only: check workbook input paths; no SEBAL execution.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --preflight
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_150039.xlsx --preflight
#
# Prepare the selected coefficients and manifest only; no SEBAL execution.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --n-realizations 200 --subset-seed 42 --prepare-only --output "../../uncertainty/bootstrap/inputs/regional/prepare_149039_200"
#
# Smoke test: one acquisition (Excel row 2), 20 realizations.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --start-row 2 --end-row 2 --n-realizations 20 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_smoke_149039"
#
# Individual workbooks: execute in separate terminals for parallel runs.
# Both commands use the same bank, count and seed, with DIFFERENT output folders.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_149039_200"
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_150039.xlsx --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_150039_200"
#
# Both workbooks sequentially in one process.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --excel InputEXCEL_LBDC_150039.xlsx --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_both_200"
#
# Restrict an individual workbook to an inclusive row range.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --start-row 10 --end-row 20 --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_149039_rows10_20"
#
# Compare the central setting with saved mean rasters before propagation.
# Replace the reference-root placeholder with your actual saved-mean folder.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --start-row 2 --end-row 2 --n-realizations 20 --subset-seed 42 --reference-root "D:/path/to/saved/mean" --output "../../uncertainty/bootstrap/inputs/regional/new_reference_check"
#
# Resume the individual run above: keep every other argument unchanged.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_149039_200" --resume
#
# Optional: use the exact preserved publication bank instead of newly generated samples.
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --coefficients "../../uncertainty/bootstrap/inputs/regional/coefficient_bank_used.csv" --n-realizations 200 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_preserved_bank_149039"
#
# Optional: propagate every coefficient pair (0 means the full input bank).
# python Run_py3_bootstrap.py --excel InputEXCEL_LBDC_149039.xlsx --n-realizations 0 --subset-seed 42 --output "../../uncertainty/bootstrap/inputs/regional/new_full_bank_149039"
